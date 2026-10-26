"""업로드된 원본 파일에서 텍스트 추출."""
import io
from pathlib import Path

from .errors import InputError

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".csv", ".json", ".html", ".htm", ".xml", ".log"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx"}


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp949"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = [f"[p.{i}]\n{page.extract_text() or ''}" for i, page in enumerate(reader.pages, 1)]
    return "\n\n".join(parts)


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        text = _pdf(data)
    elif ext == ".docx":
        text = _docx(data)
    elif ext in TEXT_EXTENSIONS:
        text = _decode(data)
    else:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise InputError(f"지원하지 않는 파일 형식입니다: {ext or '(확장자 없음)'} (지원: {supported})")
    if not text.strip():
        raise InputError("문서에서 텍스트를 추출하지 못했습니다. (스캔 PDF는 OCR이 필요합니다)")
    return text
